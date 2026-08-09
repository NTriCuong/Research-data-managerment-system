from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import NotFoundException
from app.models.enum import AuthorRole
from app.repositories.export_repository import ExportRepository

TEMPLATES_DIR = Path(__file__).resolve().parents[2] / "utils/templates_excel"
AUTHOR_PROFILE_TEMPLATE = TEMPLATES_DIR / "sheet_author_profile.xlsx"
RESEARCHES_TEMPLATE = TEMPLATES_DIR / "sheet_researches.xlsx"

MAIN_AUTHOR_ROLES = (AuthorRole.creator, AuthorRole.corresponding_author)
CONTRIBUTOR_ROLES = (AuthorRole.contributor, AuthorRole.student_member)
SUPERVISOR_ROLES = (AuthorRole.supervisor,)


def _join_authors(authors, roles: tuple[AuthorRole, ...]) -> str:
    return "; ".join(a.full_name for a in authors if a.author_role in roles)


def _format_date(value) -> str:
    return value.strftime("%d/%m/%Y") if value else ""


def _write_research_rows(sheet, researches) -> None:
    for row_offset, research in enumerate(researches):
        row = row_offset + 3
        sheet.cell(row=row, column=1, value=row_offset + 1)
        sheet.cell(row=row, column=2, value=str(research.research_id))
        sheet.cell(row=row, column=3, value=research.title)
        sheet.cell(
            row=row,
            column=4,
            value="; ".join(d.domain.domain_name for d in research.domains),
        )
        sheet.cell(row=row, column=5, value=research.department.department_name)
        sheet.cell(row=row, column=6, value=research.output_type.type_name)
        sheet.cell(row=row, column=7, value=_join_authors(research.authors, MAIN_AUTHOR_ROLES))
        sheet.cell(row=row, column=8, value=_join_authors(research.authors, CONTRIBUTOR_ROLES))
        sheet.cell(row=row, column=9, value=_join_authors(research.authors, SUPERVISOR_ROLES))
        sheet.cell(row=row, column=10, value="Đã xuất bản")
        sheet.cell(row=row, column=11, value=_format_date(research.created_at))
        sheet.cell(row=row, column=12, value=_format_date(research.approved_at))
        sheet.cell(
            row=row,
            column=13,
            value=_format_date(research.date_issued or research.approved_at),
        )


class ExportService:
    async def export_author_profile(
        self, db: AsyncSession, *, researcher_id: UUID
    ) -> tuple[BytesIO, str]:
        repo = ExportRepository(db)
        researcher = await repo.get_researcher(researcher_id)
        if researcher is None:
            raise NotFoundException("Không tìm thấy nhà nghiên cứu")

        total_researches = await repo.count_total_researches(researcher_id)
        total_pending = await repo.count_pending_researches(researcher_id)
        total_rejected = await repo.count_rejected_researches(researcher_id)
        total_published = await repo.count_published_researches(researcher_id)

        workbook = load_workbook(AUTHOR_PROFILE_TEMPLATE)
        sheet = workbook["Sheet1"]

        sheet["B2"] = researcher.full_name
        sheet["B3"] = researcher.researcher_code or ""
        sheet["B4"] = researcher.email or ""
        sheet["B5"] = researcher.department.department_name if researcher.department else ""
        sheet["B6"] = researcher.academic_title or ""
        sheet["B7"] = total_researches
        sheet["B8"] = total_pending
        sheet["B9"] = total_rejected
        sheet["B10"] = total_published
        sheet["B11"] = datetime.now().strftime("%d/%m/%Y")

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        filename = f"ho-so-{researcher.researcher_code or researcher.researcher_id}.xlsx"
        return buffer, filename

    async def export_researches_by_year(
        self, db: AsyncSession, *, year: int
    ) -> tuple[BytesIO, str]:
        repo = ExportRepository(db)
        researches = await repo.list_researches_by_year(year)

        workbook = load_workbook(RESEARCHES_TEMPLATE)
        sheet = workbook["Sheet1"]
        sheet["A1"] = f"Danh sách bài nghiên cứu năm {year}"
        _write_research_rows(sheet, researches)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        filename = f"danh-sach-bai-nghien-cuu-nam-{year}.xlsx"
        return buffer, filename

    async def export_researches_by_department(
        self, db: AsyncSession, *, department_id: UUID
    ) -> tuple[BytesIO, str]:
        repo = ExportRepository(db)
        department = await repo.get_department(department_id)
        if department is None:
            raise NotFoundException("Không tìm thấy đơn vị")

        researches = await repo.list_researches_by_department(department_id)

        workbook = load_workbook(RESEARCHES_TEMPLATE)
        sheet = workbook["Sheet1"]
        sheet["A1"] = f"Danh sách bài nghiên cứu của đơn vị {department.department_name}"
        _write_research_rows(sheet, researches)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        filename = f"danh-sach-bai-nghien-cuu-{department.department_code}.xlsx"
        return buffer, filename

    async def export_researches_by_department_and_year(
        self, db: AsyncSession, *, department_id: UUID, year: int
    ) -> tuple[BytesIO, str]:
        repo = ExportRepository(db)
        department = await repo.get_department(department_id)
        if department is None:
            raise NotFoundException("Không tìm thấy đơn vị")

        researches = await repo.list_researches_by_department_and_year(department_id, year)

        workbook = load_workbook(RESEARCHES_TEMPLATE)
        sheet = workbook["Sheet1"]
        sheet["A1"] = f"Danh sách bài nghiên cứu của đơn vị {department.department_name} năm {year}"
        _write_research_rows(sheet, researches)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        filename = f"danh-sach-bai-nghien-cuu-{department.department_code}-nam-{year}.xlsx"
        return buffer, filename

    async def export_researches_by_researcher(
        self, db: AsyncSession, *, researcher_id: UUID
    ) -> tuple[BytesIO, str]:
        repo = ExportRepository(db)
        researcher = await repo.get_researcher(researcher_id)
        if researcher is None:
            raise NotFoundException("Không tìm thấy nhà nghiên cứu")

        researches = await repo.list_researches_by_researcher(researcher_id)

        workbook = load_workbook(RESEARCHES_TEMPLATE)
        sheet = workbook["Sheet1"]
        sheet["A1"] = f"Danh sách bài nghiên cứu của {researcher.full_name}"
        _write_research_rows(sheet, researches)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        filename = f"danh-sach-bai-nghien-cuu-{researcher.researcher_code or researcher.researcher_id}.xlsx"
        return buffer, filename


export_service = ExportService()
